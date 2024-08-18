import os
import fitz  # PyMuPDF
import xml.etree.ElementTree as ET
import openpyxl  # For XLSX files
import docx  # For DOCX files
from sentence_transformers import SentenceTransformer
from pymilvus import connections, utility, FieldSchema, CollectionSchema, DataType, Collection

def extract_text_from_pdf(pdf_path):
    doc = fitz.open(pdf_path)
    text = ""
    for page_num in range(doc.page_count):
        page = doc.load_page(page_num)
        text += page.get_text()
    return text

def extract_text_from_xml(xml_path):
    tree = ET.parse(xml_path)
    root = tree.getroot()
    text = ""
    for elem in root.iter():
        if elem.text:
            text += elem.text + " "
    return text.strip()

def extract_text_from_xlsx(xlsx_path):
    workbook = openpyxl.load_workbook(xlsx_path)
    text = ""
    for sheet in workbook:
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    text += str(cell) + " "
    return text.strip()

def extract_text_from_docx(docx_path):
    doc = docx.Document(docx_path)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + " "
    return text.strip()

def chunk_text(text, max_length=500):
    # Split text into chunks of max_length
    words = text.split()
    for i in range(0, len(words), max_length):
        yield " ".join(words[i:i + max_length])

def process_files_in_directory(directory_path):
    # Initialize the model
    model = SentenceTransformer('all-MiniLM-L6-v2')

    # Connect to Milvus
    connections.connect("default", host="localhost", port="19530")

    # Verify the connection
    connected = connections.has_connection("default")
    print(f"Connected to Milvus: {connected}")

    collection_name = "document_collection"

    # List all collections
    all_collections = utility.list_collections()
    print("\nCollections:", all_collections)
    
    # Check if the collection exists
    if collection_name in all_collections:
        # Load the collection
        collection = Collection(name=collection_name)
        # Drop the collection
        collection.drop()
        print(f"Collection '{collection_name}' has been deleted.")
    else:
        print(f"Collection '{collection_name}' does not exist.")

    # Define the schema
    fields = [
        FieldSchema(name="id", dtype=DataType.INT64, is_primary=True, auto_id=True),
        FieldSchema(name="embedding", dtype=DataType.FLOAT_VECTOR, dim=384),
        FieldSchema(name="content", dtype=DataType.VARCHAR, max_length=65535)
    ]
    schema = CollectionSchema(fields, "Example collection")

    # Create a collection
    collection = Collection(collection_name, schema)

    for filename in os.listdir(directory_path):
        file_path = os.path.join(directory_path, filename)
        if filename.endswith(".pdf"):
            print(f"Processing {file_path}")
            file_text = extract_text_from_pdf(file_path)
        elif filename.endswith(".xml"):
            print(f"Processing {file_path}")
            file_text = extract_text_from_xml(file_path)
        elif filename.endswith(".xlsx"):
            print(f"Processing {file_path}")
            file_text = extract_text_from_xlsx(file_path)
        elif filename.endswith(".docx"):
            print(f"Processing {file_path}")
            file_text = extract_text_from_docx(file_path)
        else:
            continue

        # Chunk text if it's too long
        chunks = list(chunk_text(file_text))

        for chunk in chunks:
            # Convert text to vector
            text_vector = model.encode(chunk)

            # Insert vector and content into Milvus
            collection.insert([[text_vector], [chunk]])

    # Create an index for the 'embedding' field
    index_params = {
        "metric_type": "L2",
        "index_type": "IVF_FLAT",
        "params": {"nlist": 1024}
    }
    collection.create_index(field_name="embedding", index_params=index_params)

def main():
    directory_path = "embedding_documents"
    process_files_in_directory(directory_path)

if __name__ == "__main__":
    main()